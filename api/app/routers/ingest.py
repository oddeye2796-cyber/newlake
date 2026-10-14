"""데이터 넣기 (관리자용) — 원천 연결 등록·연결 시험, 파일(CSV·엑셀)·DB 표를 원천 존에 적재, 표준 존 변환.

쓰기는 모두 DB 함수(gov.register_source·ingest_records·standardize_batch)로 한다. 함수가 역할과 범위를 다시 확인한다.
외부 DB 비밀번호는 요청 하나에만 쓰고 어디에도 저장하지 않는다.
"""
import csv
import hashlib
import io
import ipaddress
import json
import re
import socket
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import PurePath
from typing import Any
from uuid import UUID

import psycopg
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from psycopg import sql
from pydantic import BaseModel, Field

from ..config import settings
from ..deps import Scoped
from .admin import admin_scoped

router = APIRouter(prefix="/api/admin", tags=["admin-ingest"])

SYSTEM_TYPES = ("ERP", "MES", "QMS", "WMS", "PLM", "SCADA", "SECS_GEM", "EDMS", "SENSOR", "FILE")
DB_KINDS = ("POSTGRESQL", "MYSQL", "MSSQL", "ORACLE", "REST", "SFTP", "FILE")
HIDDEN_TARGET_COLS = {"enterprise_id"}


# ── 원천 연결 ───────────────────────────────────────────────────────────────

@router.get("/sources")
def sources(s: Scoped = Depends(admin_scoped)):
    return s.conn.execute(
        "SELECT ss.source_system_id, ss.label_ko, ss.system_type, ss.connection_kind, ss.owner_dept,"
        "       c.db_kind, c.host, c.port, c.database_name, c.username, c.status, c.last_tested_at, c.last_error,"
        "       (SELECT count(*) FROM catalog.dataset d WHERE d.source_system_id = ss.source_system_id) AS datasets,"
        "       (SELECT max(d.last_ingested_at) FROM catalog.dataset d WHERE d.source_system_id = ss.source_system_id) AS last_ingested_at"
        "  FROM catalog.source_system ss LEFT JOIN catalog.source_connection c ON c.source_system_id = ss.source_system_id"
        " ORDER BY ss.source_system_id").fetchall()


class SourceIn(BaseModel):
    source_system_id: str = Field(pattern=r"^[A-Za-z0-9][A-Za-z0-9_-]{1,49}$")
    label: str = Field(min_length=1, max_length=100)
    system_type: str = "FILE"
    db_kind: str = "FILE"
    host: str | None = Field(default=None, max_length=255)
    port: int | None = Field(default=None, ge=1, le=65535)
    database: str | None = Field(default=None, max_length=128)
    username: str | None = Field(default=None, max_length=128)
    owner_dept: str | None = Field(default=None, max_length=100)


@router.post("/sources")
def register_source(body: SourceIn, s: Scoped = Depends(admin_scoped)):
    if body.system_type not in SYSTEM_TYPES or body.db_kind not in DB_KINDS:
        raise HTTPException(400, "원천 종류나 연결 방식이 올바르지 않습니다")
    if body.db_kind not in ("FILE", "REST", "SFTP") and not (body.host and body.database):
        raise HTTPException(400, "DB 연결에는 호스트와 데이터베이스 이름이 필요합니다")
    sid = s.conn.execute(
        "SELECT gov.register_source(%s, %s, %s, %s, %s, %s, %s, %s, NULL, %s) AS id",
        (body.source_system_id, body.label, body.system_type, body.db_kind, body.host,
         body.port or (5432 if body.db_kind == "POSTGRESQL" else None), body.database, body.username,
         body.owner_dept)).fetchone()["id"]
    return {"source_system_id": sid}


def _connection(s: Scoped, source_id: str) -> dict:
    c = s.conn.execute("SELECT * FROM catalog.source_connection WHERE source_system_id = %s", (source_id,)).fetchone()
    if c is None:
        raise HTTPException(404, f"원천 {source_id} 의 연결 정보가 없습니다")
    if c["db_kind"] != "POSTGRESQL":
        raise HTTPException(400, "지금은 PostgreSQL 만 직접 연결할 수 있습니다. 다른 DB 는 CSV·엑셀로 내보내 파일 올리기를 써 주세요")
    return c


def _check_host(host: str) -> None:
    """서버리스 시연 환경에서 내부망·루프백 주소로 연결하지 못하게 막는다 (SSRF)."""
    if settings.ingest_allow_private_hosts:
        return
    try:
        infos = socket.getaddrinfo(host, None)
    except OSError:
        raise HTTPException(400, f"호스트 {host} 를 찾을 수 없습니다")
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            raise HTTPException(400, "이 환경에서는 사내망·로컬 주소로 연결할 수 없습니다")


def _external(c: dict, password: str) -> psycopg.Connection:
    _check_host(c["host"])
    return psycopg.connect(
        host=c["host"], port=c["port"] or 5432, dbname=c["database_name"], user=c["username"], password=password,
        connect_timeout=5, sslmode="prefer", autocommit=False,
        options="-c statement_timeout=15000 -c default_transaction_read_only=on")


class Secret(BaseModel):
    password: str = Field(default="", max_length=500)


@router.post("/sources/{source_id}/test")
def test_source(source_id: str, body: Secret, s: Scoped = Depends(admin_scoped)):
    c = _connection(s, source_id)
    try:
        with _external(c, body.password) as ext:
            tables = ext.execute(
                "SELECT table_schema, table_name, table_type FROM information_schema.tables"
                " WHERE table_schema NOT IN ('pg_catalog', 'information_schema') AND table_schema NOT LIKE 'pg\\_%%'"
                " ORDER BY 1, 2 LIMIT 500").fetchall()
    except psycopg.Error as e:
        msg = str(e).strip().splitlines()[0] if str(e).strip() else type(e).__name__
        s.conn.execute("SELECT gov.record_connection_test(%s, false, %s)", (source_id, msg))
        return {"ok": False, "error": msg, "tables": []}
    s.conn.execute("SELECT gov.record_connection_test(%s, true)", (source_id,))
    return {"ok": True, "error": None,
            "tables": [{"schema": t[0], "name": t[1], "kind": "view" if t[2] == "VIEW" else "table"} for t in tables]}


class ImportIn(Secret):
    table: str = Field(pattern=r"^[^.\s]+\.[^.\s]+$", max_length=200)   # schema.table
    limit: int = Field(default=1000, ge=1, le=5000)


@router.post("/sources/{source_id}/import")
def import_table(source_id: str, body: ImportIn, s: Scoped = Depends(admin_scoped)):
    c = _connection(s, source_id)
    schema, table = body.table.split(".", 1)
    try:
        with _external(c, body.password) as ext:
            cur = ext.execute(sql.SQL("SELECT * FROM {}.{} LIMIT {}").format(
                sql.Identifier(schema), sql.Identifier(table), sql.Literal(body.limit)))
            cols = [d.name for d in cur.description or []]
            rows = [{k: _jsonable(v) for k, v in zip(cols, r)} for r in cur.fetchall()]
    except psycopg.Error as e:
        raise HTTPException(400, f"원천 DB 에서 읽지 못했습니다: {str(e).strip().splitlines()[0] if str(e).strip() else e}")
    if not rows:
        raise HTTPException(400, "표에 행이 없습니다")
    payload = json.dumps(rows, ensure_ascii=False)
    batch = s.conn.execute(
        "SELECT gov.ingest_records(%s, %s, %s::jsonb, NULL, 'table', NULL, %s, %s, NULL, %s) AS b",
        (source_id, table, payload, len(payload.encode()), hashlib.sha256(payload.encode()).hexdigest(), cols)).fetchone()["b"]
    return {"batch_id": batch, "rows": len(rows), "columns": cols}


# ── 파일 올리기 ─────────────────────────────────────────────────────────────

def _jsonable(v: Any) -> Any:
    if v is None or isinstance(v, (bool, int, float, str)):
        return v
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    if isinstance(v, Decimal):
        return str(v)
    if isinstance(v, (bytes, bytearray, memoryview)):
        return None
    return str(v)


def _headers(raw: list[Any]) -> list[str]:
    out: list[str] = []
    for i, h in enumerate(raw):
        name = str(h).strip() if h is not None else ""
        name = name or f"col_{i + 1}"
        base, n = name, 2
        while name in out:
            name, n = f"{base}_{n}", n + 1
        out.append(name)
    return out


def _parse_csv(data: bytes) -> list[dict]:
    for enc in ("utf-8-sig", "cp949"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise HTTPException(400, "파일 인코딩을 읽지 못했습니다 (UTF-8 또는 CP949 로 저장해 주세요)")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [r for r in reader if any(c.strip() for c in r)]
    if len(rows) < 2:
        raise HTTPException(400, "머리글 한 줄과 데이터가 한 줄 이상 있어야 합니다")
    head = _headers(rows[0])
    return [{h: (r[i] if i < len(r) else None) for i, h in enumerate(head)} for r in rows[1:]]


def _parse_xlsx(data: bytes, sheet: str | None) -> list[dict]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "엑셀 파일을 열지 못했습니다 (.xlsx 만 됩니다)")
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    it = (r for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r))
    first = next(it, None)
    if first is None:
        raise HTTPException(400, "시트가 비어 있습니다")
    head = _headers(list(first))
    out = [{h: _jsonable(r[i]) if i < len(r) else None for i, h in enumerate(head)} for r in it]
    if not out:
        raise HTTPException(400, "머리글 아래에 데이터가 없습니다")
    return out


@router.post("/upload")
async def upload(file: UploadFile = File(...), source_system_id: str = Form(...), entity: str | None = Form(default=None),
                 sheet: str | None = Form(default=None), s: Scoped = Depends(admin_scoped)):
    data = await file.read(settings.ingest_max_bytes + 1)
    if len(data) > settings.ingest_max_bytes:
        raise HTTPException(413, f"파일이 너무 큽니다 (최대 {settings.ingest_max_bytes // (1024 * 1024)}MB)")
    name = PurePath(file.filename or "upload.csv").name[:200]
    ext = PurePath(name).suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        rows, fmt = _parse_xlsx(data, sheet), "xlsx"
    elif ext in (".csv", ".tsv", ".txt"):
        rows, fmt = _parse_csv(data), "csv"
    else:
        raise HTTPException(400, "CSV·TSV·엑셀(.xlsx) 파일만 올릴 수 있습니다")
    if len(rows) > settings.ingest_max_rows:
        raise HTTPException(413, f"행이 너무 많습니다 (최대 {settings.ingest_max_rows:,}행)")
    ent = (entity or "").strip() or PurePath(name).stem
    batch = s.conn.execute(
        "SELECT gov.ingest_records(%s, %s, %s::jsonb, %s, %s, %s, %s, %s, NULL, %s) AS b",
        (source_system_id, ent[:100], json.dumps(rows, ensure_ascii=False), name, fmt, file.content_type,
         len(data), hashlib.sha256(data).hexdigest(), list(rows[0].keys()))).fetchone()["b"]
    return {"batch_id": batch, "rows": len(rows), "columns": list(rows[0].keys()), "entity": ent}


# ── 적재 기록·표준 변환 ─────────────────────────────────────────────────────

@router.get("/batches")
def batches(s: Scoped = Depends(admin_scoped)):
    return s.conn.execute("SELECT * FROM gov.ingest_batches(100)").fetchall()


def _norm(x: str | None) -> str:
    return re.sub(r"[^0-9a-z가-힣]", "", (x or "").lower())


def _targets(s: Scoped) -> dict[str, dict]:
    rows = s.conn.execute(
        "SELECT c.relname AS tbl, a.attname AS col, col_description(c.oid, a.attnum) AS label, a.attnotnull AS required,"
        "       format_type(a.atttypid, a.atttypmod) AS type,"
        "       EXISTS (SELECT 1 FROM pg_index i WHERE i.indrelid = c.oid AND i.indisprimary AND a.attnum = ANY (i.indkey)) AS pk,"
        "       a.atthasdef AS has_default"
        "  FROM pg_class c JOIN pg_namespace n ON n.oid = c.relnamespace AND n.nspname = 'ksx'"
        "  JOIN pg_attribute a ON a.attrelid = c.oid AND a.attnum > 0 AND NOT a.attisdropped"
        " WHERE c.relkind IN ('r', 'p') AND NOT c.relispartition"
        " ORDER BY c.relname, a.attnum").fetchall()
    classes = {r["physical_table"]: r for r in s.conn.execute(
        "SELECT DISTINCT ON (physical_table) physical_table, label_ko, local_name, class_iri FROM onto.class"
        " WHERE physical_table IS NOT NULL ORDER BY physical_table, is_standard DESC").fetchall()}
    out: dict[str, dict] = {}
    for r in rows:
        if r["col"] in HIDDEN_TARGET_COLS or r["col"].startswith("lake_"):
            continue
        cl = classes.get(r["tbl"])
        t = out.setdefault(r["tbl"], {"table": r["tbl"], "label_ko": cl["label_ko"] if cl else None,
                                      "class_name": cl["local_name"] if cl else None, "columns": []})
        t["columns"].append({"name": r["col"], "label_ko": r["label"], "type": r["type"],
                             "required": r["required"] and not r["has_default"], "pk": r["pk"]})
    return out


@router.get("/targets")
def targets(s: Scoped = Depends(admin_scoped)):
    return sorted(_targets(s).values(), key=lambda t: (t["label_ko"] is None, t["table"]))


def suggest(source_cols: list[str], entity: str | None, targets_: dict[str, dict]) -> tuple[str | None, dict[str, str]]:
    """열 이름·한글 설명이 같은 것끼리 짝짓고, 가장 많이 맞는 표를 고른다."""
    best, best_score, best_map = None, 0.0, {}
    ent = _norm(entity)
    for name, t in targets_.items():
        index: dict[str, str] = {}
        for c in t["columns"]:
            for key in (_norm(c["name"]), _norm(c["label_ko"])):
                if key:
                    index.setdefault(key, c["name"])
        m: dict[str, str] = {}
        for sc in source_cols:
            tc = index.get(_norm(sc))
            if tc and tc not in m.values():
                m[sc] = tc
        score = len(m) + (1.5 if ent and (_norm(name) in ent or (t["label_ko"] and _norm(t["label_ko"]) in ent)) else 0)
        if score > best_score:
            best, best_score, best_map = name, score, m
    return best, best_map


@router.get("/batches/{batch_id}")
def batch(batch_id: UUID, table: str | None = None, s: Scoped = Depends(admin_scoped)):
    info = s.conn.execute("SELECT * FROM gov.ingest_batches(500) WHERE batch_id = %s", (batch_id,)).fetchone()
    if info is None:
        raise HTTPException(404, "배치가 없거나 볼 수 있는 범위 밖입니다")
    preview = s.conn.execute("SELECT * FROM gov.ingest_batch_preview(%s, 20)", (batch_id,)).fetchall()
    cols: list[str] = list(info["columns"] or [])
    for p in preview:
        for k in p["payload"]:
            if k not in cols:
                cols.append(k)
    tg = _targets(s)
    suggested, mapping = suggest(cols, info["source_entity"], tg)
    if table and table in tg:
        _, mapping = suggest(cols, None, {table: tg[table]})
        suggested = table
    return {"batch": info, "columns": cols, "rows": [p["payload"] for p in preview],
            "suggested_table": suggested, "suggested_map": mapping}


class StandardizeIn(BaseModel):
    table: str = Field(pattern=r"^[a-z][a-z0-9_]{0,62}$")
    column_map: dict[str, str]
    site_id: str | None = Field(default=None, max_length=50)
    plant_id: str | None = Field(default=None, max_length=50)


@router.post("/batches/{batch_id}/standardize")
def standardize(batch_id: UUID, body: StandardizeIn, s: Scoped = Depends(admin_scoped)):
    return s.conn.execute(
        "SELECT gov.standardize_batch(%s, %s, %s::jsonb, %s, %s) AS r",
        (batch_id, body.table, json.dumps({k: v for k, v in body.column_map.items() if v}),
         body.site_id or None, body.plant_id or None)).fetchone()["r"]
