import io
from datetime import date
import uuid

from openpyxl import Workbook

from .conftest import as_user

ADMIN = as_user("u-choi")


def _source(client) -> str:
    sid = "T-" + uuid.uuid4().hex[:8].upper()
    r = client.post("/api/admin/sources", headers=ADMIN,
                    json={"source_system_id": sid, "label": "테스트 원천", "system_type": "MES", "db_kind": "FILE"})
    assert r.status_code == 200, r.text
    return sid


def test_ingest_requires_ingest_role(client):
    assert client.get("/api/admin/batches", headers=as_user("u-jung")).status_code == 403


def test_register_and_list_source(client):
    sid = _source(client)
    rows = {s["source_system_id"]: s for s in client.get("/api/admin/sources", headers=ADMIN).json()}
    assert rows[sid]["db_kind"] == "FILE" and rows[sid]["label_ko"] == "테스트 원천"
    assert "MES-A" in rows


def test_csv_upload_suggest_and_standardize(client):
    sid = _source(client)
    lot = "LOT-T" + uuid.uuid4().hex[:6].upper()
    body = f"로트식별자,로트명,PLANT_ID,lot_created_quantity\n{lot},테스트,P1,12\n{lot}-B,잘못된 수량,P1,abc\n".encode("cp949")
    r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid, "entity": "lot_master"},
                    files={"file": ("lots.csv", body, "text/csv")})
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["rows"] == 2 and b["columns"][0] == "로트식별자"

    d = client.get(f"/api/admin/batches/{b['batch_id']}", headers=ADMIN).json()
    assert d["suggested_table"] == "lot"
    assert d["suggested_map"]["로트식별자"] == "lot_id" and d["suggested_map"]["PLANT_ID"] == "plant_id"

    r = client.post(f"/api/admin/batches/{b['batch_id']}/standardize", headers=ADMIN,
                    json={"table": "lot", "column_map": d["suggested_map"], "site_id": "HQ"})
    assert r.status_code == 200, r.text
    res = r.json()
    assert (res["rows_in"], res["rows_out"], res["rows_rejected"]) == (2, 1, 1)
    assert res["std_dataset_id"] == "std.lot"
    assert client.get(f"/api/lots/{lot}", headers=ADMIN).status_code == 200

    hist = {x["batch_id"]: x for x in client.get("/api/admin/batches", headers=ADMIN).json()}
    assert hist[b["batch_id"]]["status"] == "STANDARDIZED" and hist[b["batch_id"]]["std_rows_out"] == 1
    ds = client.get(f"/api/admin/datasets/{hist[b['batch_id']]['dataset_id']}", headers=ADMIN).json()
    assert ds["downstream"][0]["dataset_id"] == "std.lot"


def test_xlsx_upload(client):
    sid = _source(client)
    wb = Workbook()
    ws = wb.active
    ws.append(["설비ID", "설비명"])
    ws.append(["EQ-X1", "테스트 설비"])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid},
                    files={"file": ("equipment.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    assert r.json()["rows"] == 1 and r.json()["entity"] == "equipment"


def test_standardize_rejects_missing_pk(client):
    sid = _source(client)
    r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid},
                    files={"file": ("x.csv", b"name\nfoo\n", "text/csv")})
    r = client.post(f"/api/admin/batches/{r.json()['batch_id']}/standardize", headers=ADMIN,
                    json={"table": "lot", "column_map": {"name": "lot_name"}})
    assert r.status_code == 409 and "기본키" in r.json()["detail"]


def test_upload_to_unknown_source_is_refused(client):
    r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": "NOPE-404"},
                    files={"file": ("x.csv", b"a\n1\n", "text/csv")})
    assert r.status_code == 409


def test_db_test_only_postgres(client):
    sid = "T-" + uuid.uuid4().hex[:8].upper()
    client.post("/api/admin/sources", headers=ADMIN, json={"source_system_id": sid, "label": "MySQL 원천",
                "system_type": "ERP", "db_kind": "MYSQL", "host": "db.example", "database": "erp", "username": "r"})
    assert client.post(f"/api/admin/sources/{sid}/test", headers=ADMIN, json={"password": "x"}).status_code == 400


def test_db_connect_and_import(client):
    # 원천 DB 대신 레이크 DB 자체에 svc_app 로 붙어 본다 (읽기 전용 트랜잭션)
    from app.config import settings
    from psycopg.conninfo import conninfo_to_dict

    ci = conninfo_to_dict(settings.database_url)
    sid = "T-" + uuid.uuid4().hex[:8].upper()
    r = client.post("/api/admin/sources", headers=ADMIN, json={
        "source_system_id": sid, "label": "레이크 자신", "system_type": "MES", "db_kind": "POSTGRESQL",
        "host": ci.get("host", "localhost"), "port": int(ci.get("port", 5432)), "database": ci["dbname"], "username": ci["user"]})
    assert r.status_code == 200, r.text
    t = client.post(f"/api/admin/sources/{sid}/test", headers=ADMIN, json={"password": ci.get("password", "")}).json()
    assert t["ok"], t
    assert any(x["schema"] == "ksx" and x["name"] == "lot" for x in t["tables"])
    r = client.post(f"/api/admin/sources/{sid}/import", headers=ADMIN,
                    json={"password": ci.get("password", ""), "table": "onto.class", "limit": 5})
    assert r.status_code == 200, r.text
    assert r.json()["rows"] == 5
    src = {s["source_system_id"]: s for s in client.get("/api/admin/sources", headers=ADMIN).json()}[sid]
    assert src["status"] == "OK" and src["datasets"] == 1


def test_scm_source_type(client):
    sid = "T-" + uuid.uuid4().hex[:8].upper()
    r = client.post("/api/admin/sources", headers=ADMIN,
                    json={"source_system_id": sid, "label": "SCM 원천", "system_type": "SCM", "db_kind": "FILE"})
    assert r.status_code == 200, r.text


def test_auto_map_purchase_order_file(client):
    """현장 발주·입고 엑셀처럼 열 이름이 표준과 다른 파일도 표와 열을 한 번에 제안한다."""
    sid = _source(client)
    tag = uuid.uuid4().hex[:6].upper()
    wb = Workbook()
    ws = wb.active
    ws.append(["진행상태", "납품요청일", "발주번호", "구분", "업체명", "품번", "품명", "두께", "폭", "발주수량", "가입고수량", "미입고수량"])
    ws.append(["입고", date(2026, 10, 1), f"SC{tag}01", "코일류", "(주)두원에스티에스", f"S{tag}1", "STS304", 0.4, 285.0, 1000, 1078, -78])
    ws.append(["발주", date(2026, 10, 2), f"SC{tag}02", "코일류", "(주)대한", f"S{tag}2", "STS430", 0.5, 300.0, 500, 0, 500])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid},
                    files={"file": ("owp_204010_01.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    bid = r.json()["batch_id"]

    d = client.get(f"/api/admin/batches/{bid}", headers=ADMIN).json()
    assert d["suggested_table"] == "purchase_order_line"
    m, conf = d["suggested_map"], d["confidence"]
    assert m["발주번호"] == "purchase_document_number" and conf["발주번호"] >= 0.9   # 발주번호는 문서 번호, 식별자는 "문서번호-순번" 자동
    assert m["품번"] == "item_id" and m["품명"] == "item_name" and m["발주수량"] == "purchase_order_quantity"
    assert m["납품요청일"] == "expected_delivery_date"
    assert "업체명" not in m  # 이름을 식별자 열에 넣지 않는다
    assert d["candidates"][0]["table"] == "purchase_order_line" and d["candidates"][0]["pk_ok"]
    assert d["profiles"]["납품요청일"] == "date" and d["profiles"]["발주수량"] == "num"

    # 품번이 품목 마스터에 없으면 거부된다
    body = {"table": "purchase_order_line", "column_map": m, "site_id": "HQ", "plant_id": "P1"}
    res = client.post(f"/api/admin/batches/{bid}/standardize", headers=ADMIN, json=body).json()
    assert res["rows_out"] == 0 and res["rows_rejected"] == 2

    # 마스터 자동 등록을 켜면 품목을 먼저 만들고 넣는다
    res = client.post(f"/api/admin/batches/{bid}/standardize", headers=ADMIN, json=body | {"create_refs": True}).json()
    assert (res["rows_out"], res["rows_rejected"]) == (2, 0), res
    assert res["refs_created"] == {"ksx.item_master": 2}


def test_source_crud(client):
    sid = _source(client)
    # 같은 ID 로 다시 등록하면 덮어쓰지 않고 거절
    dup = client.post("/api/admin/sources", headers=ADMIN, json={"source_system_id": sid, "label": "x", "system_type": "MES", "db_kind": "FILE"})
    assert dup.status_code == 409
    r = client.put(f"/api/admin/sources/{sid}", headers=ADMIN,
                   json={"source_system_id": sid, "label": "이름 바꿈", "system_type": "SCM", "db_kind": "FILE", "owner_dept": "구매팀"})
    assert r.status_code == 200, r.text
    row = {s["source_system_id"]: s for s in client.get("/api/admin/sources", headers=ADMIN).json()}[sid]
    assert (row["label_ko"], row["system_type"], row["owner_dept"]) == ("이름 바꿈", "SCM", "구매팀")
    assert client.put("/api/admin/sources/NOPE-404", headers=ADMIN,
                      json={"source_system_id": "NOPE-404", "label": "x", "db_kind": "FILE"}).status_code == 404
    assert client.delete(f"/api/admin/sources/{sid}", headers=as_user("u-jung")).status_code == 403

    # 데이터가 있으면 함께 지우기를 골라야 지워진다
    up = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid},
                     files={"file": ("x.csv", b"a,b\n1,2\n3,4\n", "text/csv")}).json()
    r = client.delete(f"/api/admin/sources/{sid}", headers=ADMIN)
    assert r.status_code == 409 and "함께 지우기" in r.json()["detail"]
    r = client.delete(f"/api/admin/sources/{sid}?purge=true", headers=ADMIN)
    assert r.status_code == 200, r.text
    assert r.json() == {"source_system_id": sid, "datasets": 1, "batches": 1, "records": 2}
    assert sid not in {s["source_system_id"] for s in client.get("/api/admin/sources", headers=ADMIN).json()}
    assert client.get(f"/api/admin/batches/{up['batch_id']}", headers=ADMIN).status_code == 404

    # 표준 존에 쓰인 시드 원천은 지울 수 없다
    assert client.delete("/api/admin/sources/MES-A?purge=true", headers=ADMIN).status_code == 409


def test_upload_sniffs_real_format_not_extension(client):
    """MES·ERP 가 .xlsx 라는 이름으로 내보내는 옛 .xls·HTML 표·XML 스프레드시트도 내용으로 가려 읽는다."""
    import xlwt

    sid = _source(client)
    html = ("<html><body><table><tr><th>LOT</th><th>수량</th></tr>"
            "<tr><td>H001</td><td>10</td></tr><tr><td>H002</td><td>20</td></tr></table></body></html>").encode("cp949")
    r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid},
                    files={"file": ("export.xlsx", html, "application/vnd.ms-excel")})
    assert r.status_code == 200, r.text
    assert r.json()["rows"] == 2 and r.json()["columns"] == ["LOT", "수량"]

    xml = ('<?xml version="1.0"?><Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"'
           ' xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet ss:Name="S1"><Table>'
           '<Row><Cell><Data ss:Type="String">LOT</Data></Cell><Cell><Data ss:Type="String">수량</Data></Cell></Row>'
           '<Row><Cell><Data ss:Type="String">X001</Data></Cell><Cell ss:Index="2"><Data ss:Type="Number">5</Data></Cell></Row>'
           '</Table></Worksheet></Workbook>').encode("utf-8")
    r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid},
                    files={"file": ("export2.xlsx", xml, "application/vnd.ms-excel")})
    assert r.status_code == 200, r.text
    assert r.json()["rows"] == 1 and r.json()["columns"] == ["LOT", "수량"]

    wb = xlwt.Workbook()
    ws = wb.add_sheet("S")
    for j, h in enumerate(["LOT", "수량", "일자"]):
        ws.write(0, j, h)
    ws.write(1, 0, "B001"); ws.write(1, 1, 7); ws.write(1, 2, "2026-10-01")
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid},
                    files={"file": ("old.xlsx", buf.getvalue(), "application/vnd.ms-excel")})
    assert r.status_code == 200, r.text
    assert r.json()["rows"] == 1 and r.json()["columns"] == ["LOT", "수량", "일자"]

    r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid},
                    files={"file": ("junk.xlsx", b"\x00\x01\x02 not a spreadsheet", "application/octet-stream")})
    assert r.status_code == 400 and "엑셀 파일이 아닌" in r.json()["detail"]


def test_upload_repairs_broken_xlsx_styles(client):
    """MES 가 만든 .xlsx 의 styles.xml 이 규격에 어긋나도(openpyxl: could not read stylesheet) 보정해서 읽는다. 날짜 서식은 지킨다."""
    import re
    import zipfile
    from datetime import date

    from openpyxl import Workbook

    sid = _source(client)
    wb = Workbook()
    ws = wb.active
    ws.append(["LOT", "수량", "일자"])
    ws.append(["R001", 4, date(2026, 10, 1)])
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))

    def rebuild(mod):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            for n in src.namelist():
                d = src.read(n)
                z.writestr(n, mod(d) if n == "xl/styles.xml" else d)
        return out.getvalue()

    bad_color = rebuild(lambda d: re.sub(rb'<fonts count="1">', rb'<fonts count="2"><font><color rgb="ZZ"/></font>', d))
    bad_xml = rebuild(lambda d: d.replace(b"</styleSheet>", b""))
    for name, data in (("bad_color.xlsx", bad_color), ("bad_xml.xlsx", bad_xml)):
        r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid},
                        files={"file": (name, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert r.status_code == 200, (name, r.text)
        body = r.json()
        assert body["rows"] == 1 and body["columns"] == ["LOT", "수량", "일자"]
        pv = client.get(f"/api/admin/batches/{body['batch_id']}", headers=ADMIN).json()
        assert pv["rows"][0]["일자"].startswith("2026-10-01"), (name, pv["rows"][0])


FIXTURES = __import__("pathlib").Path(__file__).parent / "fixtures"


def _upload(client, sid: str, name: str) -> dict:
    data = (FIXTURES / name).read_bytes()
    r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid},
                    files={"file": (name, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, (name, r.text)
    return r.json()


def test_real_mes_export_maps_to_production_performance(client):
    """실제 MES(TreeGrid) 실적 내보내기와 같은 구조의 파일(규격에 어긋난 styles.xml, 줄바꿈 든 머리글 43개).
    열리고, 생산실적 표로 자동 대응되며, 원천에 없는 생산실적순번은 1 로 채워져 표준 존에 들어간다."""
    import zipfile

    from openpyxl import load_workbook

    raw = (FIXTURES / "mes_treegrid_export.xlsx").read_bytes()
    try:                                   # openpyxl 은 이 파일을 그대로는 못 연다 — 보정 경로를 타는지 확인
        load_workbook(io.BytesIO(raw), read_only=True)
        assert False, "fixture no longer exercises the styles repair path"
    except ValueError:
        pass
    assert "xl/drawings/drawing1.xml" not in zipfile.ZipFile(io.BytesIO(raw)).namelist()

    sid = _source(client)
    up = _upload(client, sid, "mes_treegrid_export.xlsx")
    assert up["rows"] == 3 and len(up["columns"]) == 43 and up["columns"][0] == "배정\n번호"

    pv = client.get(f"/api/admin/batches/{up['batch_id']}", headers=ADMIN).json()
    assert pv["suggested_table"] == "production_performance", pv["candidates"]
    m = pv["suggested_map"]
    assert m["배정\n번호"] == "production_performance_id" and m["작업일"] == "production_performance_date"
    assert m["양품\n수량"] == "produced_quantity" and m["불량\n수량"] == "defect_quantity" and m["품번"] == "item_id"
    assert m["시작시간"] == "execution_start_date_time" and m["종료시간"] == "execution_end_date_time"
    assert m["작업자"] == "personnel_id" and "production_performance_sequence" not in m.values()
    assert pv["rows"][0]["작업일"].startswith("2026-09-30")

    assert pv["auto_keys"] == {"seq": {"production_performance_sequence": ["production_performance_id"]}, "derived": {}}
    tg = {t["table"]: t for t in client.get("/api/admin/targets", headers=ADMIN).json()}
    assert tg["operation_start"]["unique_keys"] == [["lot_id", "operation_id", "operation_start_time"]]   # 이벤트 식별자의 자연 키
    pol = tg["purchase_order_line"]   # 라인 표에는 헤더 열(공급업체 등)도 대응할 수 있다
    assert pol["header"] == "purchase_order_header" and any(c["name"] == "vendor_id" and c.get("header") for c in pol["columns"])
    assert not any(c["name"].endswith("_sequence_no") for t in tg.values() for c in t["columns"])

    r = client.post(f"/api/admin/batches/{up['batch_id']}/standardize", headers=ADMIN,
                    json={"table": "production_performance", "column_map": m, "site_id": "HQ", "plant_id": "P1",
                          "create_refs": True, "keep_extras": True, "auto_quality": True})
    assert r.status_code == 200, r.text
    res = r.json()
    assert (res["rows_in"], res["rows_out"], res["rows_rejected"]) == (3, 3, 0), res["errors"]
    assert isinstance(res["extensions"], dict) and "matched" in res["extensions"]   # 대응 안 된 35열은 확장 속성으로 보관
    ea = client.get("/api/admin/extension-attributes?table=production_performance&limit=100", headers=ADMIN).json()
    kept = {x["attr_name"] for x in ea["summary"]}
    assert {"설비명", "공정명", "Shift"} <= kept, kept


def test_real_scm_export_maps_to_purchase_order(client):
    """실제 SCM 발주 조회 내보내기와 같은 열(진행상태·구분·업체명·두께·폭·미입고수량 포함)."""
    sid = _source(client)
    up = _upload(client, sid, "scm_po_export.xlsx")
    assert up["rows"] == 3 and up["columns"][:3] == ["진행상태", "납품요청일", "발주번호"]
    pv = client.get(f"/api/admin/batches/{up['batch_id']}", headers=ADMIN).json()
    assert pv["suggested_table"] == "purchase_order_line", pv["candidates"]
    m = pv["suggested_map"]
    assert m["발주번호"] == "purchase_document_number" and m["납품요청일"] == "expected_delivery_date"
    assert pv["auto_keys"] == {"seq": {"purchase_document_sequence": ["purchase_document_number"]}, "derived": {}}
    assert m["발주수량"] == "purchase_order_quantity" and m["가입고수량"] == "purchase_goods_receipt_quantity"
    assert set(pv["columns"]) - set(m) == {"구분", "업체명", "두께", "폭", "미입고수량"}
    r = client.post(f"/api/admin/batches/{up['batch_id']}/standardize", headers=ADMIN,
                    json={"table": "purchase_order_line", "column_map": m, "site_id": "HQ", "plant_id": "P1",
                          "create_refs": True, "keep_extras": True})
    assert r.status_code == 200, r.text
    assert (r.json()["rows_out"], r.json()["rows_rejected"]) == (3, 0), r.json()["errors"]


def test_po_lines_keep_every_line(client):
    """발주 한 건에 품목 줄이 여럿인 라인 단위 파일: 줄이 겹쳐 사라지지 않는다.
    발주번호 → 구매문서번호, 순번은 문서별 1,2,3…, 발주 헤더는 함께 채워진다 (표준 v1.0 헤더-라인)."""
    sid = _source(client)
    po = "PO" + uuid.uuid4().hex[:6].upper()
    body = (f"발주번호,품번,발주수량,납품요청일,업체명\n{po},A-1,10,2026-10-01,(주)두원\n{po},A-2,20,2026-10-01,(주)두원\n"
            f"{po},A-3,30,2026-10-02,(주)두원\n{po}X,A-1,5,2026-10-03,(주)대한\n").encode("utf-8")
    r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid, "entity": "purchase_order"},
                    files={"file": ("po_lines.csv", body, "text/csv")})
    assert r.status_code == 200, r.text
    bid = r.json()["batch_id"]
    pv = client.get(f"/api/admin/batches/{bid}", headers=ADMIN).json()
    assert pv["suggested_table"] == "purchase_order_line" and pv["suggested_map"]["발주번호"] == "purchase_document_number"
    assert pv["dup_keys"] == 0 and "purchase_document_sequence" in pv["auto_keys"]["seq"]
    r = client.post(f"/api/admin/batches/{bid}/standardize", headers=ADMIN,
                    json={"table": "purchase_order_line", "column_map": pv["suggested_map"], "site_id": "HQ", "plant_id": "P1",
                          "create_refs": True, "keep_extras": True, "auto_quality": False})
    assert r.status_code == 200, r.text
    assert (r.json()["rows_out"], r.json()["rows_rejected"]) == (4, 0), r.json()["errors"]
    ea = client.get("/api/admin/extension-attributes?table=purchase_order_line&limit=200", headers=ADMIN).json()
    keys = {x["record_key"] for x in ea["rows"] if x["attr_name"] == "업체명" and x["record_key"].startswith(po)}
    assert keys == {f"{po}|1", f"{po}|2", f"{po}|3", f"{po}X|1"}, keys
    # 같은 파일을 다시 넣어도 같은 식별자로 덮어쓴다 (줄이 늘지 않는다)
    r = client.post(f"/api/admin/batches/{bid}/standardize", headers=ADMIN,
                    json={"table": "purchase_order_line", "column_map": pv["suggested_map"], "site_id": "HQ", "plant_id": "P1", "keep_extras": False})
    assert r.status_code == 200 and r.json()["rows_out"] == 4


def test_blank_key_rows_warned_and_rejected_readably(client):
    """합계 줄처럼 키가 빈 행: 미리보기가 미리 세어 주고, 넣을 때 거부 사유가 사람 말로 나온다."""
    sid = _source(client)
    po = "PO" + uuid.uuid4().hex[:6].upper()
    body = (f"발주번호,품번,발주수량,납품요청일\n{po},A-1,10,2026-10-01\n{po},A-2,20,2026-10-01\n,합계,30,\n").encode("utf-8")
    r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid, "entity": "purchase_order"},
                    files={"file": ("po_totals.csv", body, "text/csv")})
    assert r.status_code == 200, r.text
    bid = r.json()["batch_id"]
    pv = client.get(f"/api/admin/batches/{bid}", headers=ADMIN).json()
    assert pv["suggested_table"] == "purchase_order_line"
    bk = pv["blank_keys"]
    assert bk["rows"] == 1 and bk["first"] == 3 and "발주번호" in bk["columns"], bk
    r = client.post(f"/api/admin/batches/{bid}/standardize", headers=ADMIN,
                    json={"table": "purchase_order_line", "column_map": pv["suggested_map"], "site_id": "HQ", "plant_id": "P1",
                          "create_refs": True, "keep_extras": False, "auto_quality": False})
    assert r.status_code == 200, r.text
    res = r.json()
    assert (res["rows_out"], res["rows_rejected"]) == (2, 1), res["errors"]
    err = res["errors"][0]
    assert err["record_seq"] == 3
    assert "키 열" in err["error"] and "비어" in err["error"] and "발주번호" in err["error"], err
    assert err["raw"] and "null value" in err["raw"], err


def test_line_file_fills_header_and_rollback_clears_it(client):
    """표준 v1.0 헤더-라인: 발주 라인 파일에 헤더 열(공급업체)이 함께 오면 발주 헤더를 같이 채우고, 되돌리면 같이 지운다."""
    sid = _source(client)
    po = "PO" + uuid.uuid4().hex[:6].upper()
    body = (f"발주번호,품번,발주수량,공급업체\n{po},A-1,10,V-01\n{po},A-2,20,V-01\n{po}X,A-1,5,V-02\n").encode("utf-8")
    r = client.post("/api/admin/upload", headers=ADMIN, data={"source_system_id": sid, "entity": "purchase_order"},
                    files={"file": ("po_hdr.csv", body, "text/csv")})
    bid = r.json()["batch_id"]
    m = {"발주번호": "purchase_document_number", "품번": "item_id", "발주수량": "purchase_order_quantity", "공급업체": "vendor_id"}
    res = client.post(f"/api/admin/batches/{bid}/standardize", headers=ADMIN,
                      json={"table": "purchase_order_line", "column_map": m, "site_id": "HQ", "plant_id": "P1",
                            "create_refs": True, "keep_extras": False, "auto_quality": False}).json()
    assert (res["rows_out"], res["rows_rejected"]) == (3, 0), res["errors"]
    assert res["header"] == "ksx.purchase_order_header" and res["header_rows"] == 2, res
    rb = client.post(f"/api/admin/batches/{bid}/rollback", headers=ADMIN).json()
    assert rb["rows"] == 3 and rb["header_rows"] == 2, rb
